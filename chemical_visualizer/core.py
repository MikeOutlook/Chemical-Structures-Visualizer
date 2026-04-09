from dataclasses import dataclass
from pathlib import Path
import re
from typing import Callable, List, Optional

import openpyxl
import pandas as pd
from openpyxl.drawing.image import Image as OpenPyXLImage
from openpyxl.styles import Font
from rdkit import Chem
from rdkit.Chem import Draw


# 这个模块集中处理数据读取、结构图生成和 Excel 导出。
ProgressCallback = Callable[[int, int], None]


@dataclass
class CompoundRecord:
    """Represents one compound entry loaded from CSV or manual input."""

    # `status` / `image_path` / `error` 会随着处理流程不断更新。
    index: object
    smiles: str
    status: str = "pending"
    image_path: Optional[str] = None
    error: Optional[str] = None


class CompoundProcessor:
    """Loads compounds, renders 2D structure images, and tracks results."""

    def __init__(self, image_size=(300, 200)):
        # 在初始化时统一校验尺寸，避免后续每次生成图片都重复检查。
        self.image_size = self._normalize_image_size(image_size)
        self.compounds = []  # type: List[CompoundRecord]

    @staticmethod
    def _normalize_image_size(image_size):
        # 所有图片尺寸都走同一套入口，保证 CLI、GUI 和测试行为一致。
        width, height = image_size
        if width <= 0 or height <= 0:
            raise ValueError("Image size must use positive integers")
        return int(width), int(height)

    @staticmethod
    def _build_image_filename(index):
        # 文件名只保留安全字符，避免索引里带空格或符号时写文件失败。
        safe_index = re.sub(r"[^A-Za-z0-9_.-]+", "_", str(index)).strip("_")
        return "compound_{0}.png".format(safe_index or "item")

    def load_csv(self, filepath):
        """Load compounds from a CSV file."""

        dataframe = pd.read_csv(filepath)
        if "SMILES" not in dataframe.columns:
            raise ValueError("CSV must contain a 'SMILES' column")

        # `Index` 列是可选的；没有时就按行号从 1 开始自动编号。
        has_index_column = "Index" in dataframe.columns
        self.compounds = []

        for row_number, row in dataframe.iterrows():
            smiles = str(row["SMILES"]).strip()
            index = row["Index"] if has_index_column else row_number + 1
            self.compounds.append(CompoundRecord(index=index, smiles=smiles))

        return len(self.compounds)

    def add_smiles(self, smiles, index=None):
        """Add a single SMILES string after validating it with RDKit."""

        normalized_smiles = smiles.strip()
        if not normalized_smiles:
            return False

        # 先尝试解析，只有合法 SMILES 才进入列表。
        molecule = Chem.MolFromSmiles(normalized_smiles)
        if molecule is None:
            return False

        compound_index = index if index is not None else len(self.compounds) + 1
        self.compounds.append(CompoundRecord(index=compound_index, smiles=normalized_smiles))
        return True

    def process_compound(self, compound, output_dir):
        """Render a single compound image and update its processing status."""

        molecule = Chem.MolFromSmiles(compound.smiles)
        if molecule is None:
            # 错误信息直接写回记录，方便 GUI 列表和导出逻辑读取。
            compound.status = "error"
            compound.error = "Invalid SMILES"
            compound.image_path = None
            return False

        output_dir_path = Path(output_dir)
        # 输出目录不存在时自动创建，调用方无需提前准备。
        output_dir_path.mkdir(parents=True, exist_ok=True)

        image = Draw.MolToImage(molecule, size=self.image_size)
        image_path = output_dir_path / self._build_image_filename(compound.index)
        image.save(image_path)

        # 成功后记录图片路径，后续预览和 Excel 嵌图都会复用。
        compound.status = "success"
        compound.error = None
        compound.image_path = str(image_path)
        return True

    def process_all(self, output_dir, progress_callback=None, only_pending=False):
        """Process all compounds or only those still marked as pending."""

        targets = self.compounds
        if only_pending:
            # 重复导出时只补跑待处理项，避免重复渲染相同结构图。
            targets = [compound for compound in self.compounds if compound.status == "pending"]

        success_count = 0
        fail_count = 0
        total = len(targets)

        for current, compound in enumerate(targets, start=1):
            if self.process_compound(compound, output_dir):
                success_count += 1
            else:
                fail_count += 1

            if progress_callback is not None:
                # 进度展示由上层决定，这里只负责通知当前进度。
                progress_callback(current, total)

        return success_count, fail_count

    def clear(self):
        """Remove all loaded compounds."""

        self.compounds = []

    def get_compound(self, index):
        """Return the compound with the requested index, if present."""

        for compound in self.compounds:
            if compound.index == index:
                return compound
        return None

    def generate_preview_image(self, smiles, size=(400, 300)):
        """Generate a preview image for a SMILES string."""

        width, height = self._normalize_image_size(size)
        molecule = Chem.MolFromSmiles(smiles)
        if molecule is None:
            return None
        return Draw.MolToImage(molecule, size=(width, height))


def create_excel_with_images(processor, output_path, image_dir=None):
    """Create an Excel workbook that embeds generated structure images."""

    # 保留旧参数名以兼容已有调用方，当前实现本身不再使用它。
    del image_dir  # Retained for backward compatibility with existing callers.

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Chemical Structures"
    worksheet.freeze_panes = "A2"

    headers = ("Index", "SMILES", "Structure")
    for column, value in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=column, value=value)
        cell.font = Font(bold=True)

    for row_number, compound in enumerate(processor.compounds, start=2):
        worksheet.cell(row=row_number, column=1, value=compound.index)
        worksheet.cell(row=row_number, column=2, value=compound.smiles)

        if compound.status == "success" and compound.image_path and Path(compound.image_path).exists():
            # 只有文件真实存在时才插图，避免生成损坏的工作簿。
            image = OpenPyXLImage(compound.image_path)
            image.width = 150
            image.height = 100
            worksheet.add_image(image, "C{0}".format(row_number))
            worksheet.row_dimensions[row_number].height = 80

    worksheet.column_dimensions["A"].width = 12
    worksheet.column_dimensions["B"].width = 60
    worksheet.column_dimensions["C"].width = 24

    output_path = Path(output_path)
    # 允许直接导出到尚不存在的目录。
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
